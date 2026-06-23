"""Excel (XLSX) formatting for human review of enriched output.

:func:`write_formatted_xlsx` writes a review-friendly sheet: a frozen, filtered
header; autofit-ish column widths; highlighted ``needs_review`` rows;
color-coded ``match_confidence``; and clickable website / maps URLs.

Styling is keyed by column name and degrades gracefully — columns that aren't
present are simply skipped — so it is safe for any DataFrame.
"""
from __future__ import annotations

from pathlib import Path
from typing import Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PathLike = Union[str, Path]

_HEADER_FILL = PatternFill("solid", fgColor="366092")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_REVIEW_FILL = PatternFill("solid", fgColor="FCE4E4")  # light pink row highlight
_LINK_FONT = Font(color="0563C1", underline="single")

# Obvious, distinct styles per confidence level (Excel's status palette).
_CONFIDENCE_STYLE = {
    "high":   (PatternFill("solid", fgColor="C6EFCE"), Font(bold=True, color="006100")),
    "medium": (PatternFill("solid", fgColor="FFEB9C"), Font(bold=True, color="9C6500")),
    "low":    (PatternFill("solid", fgColor="FFCC99"), Font(bold=True, color="9C4500")),
    "none":   (PatternFill("solid", fgColor="FFC7CE"), Font(bold=True, color="9C0006")),
}

# URL columns to render as clickable hyperlinks.
_URL_COLUMNS = ("google_website", "google_maps_uri", "official_website_candidate")

_MIN_WIDTH = 10
_MAX_WIDTH = 60


def write_formatted_xlsx(df: pd.DataFrame, path: PathLike) -> Path:
    """Write ``df`` to ``path`` as a review-formatted XLSX. Returns the path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    safe = df.fillna("")
    safe.to_excel(path, index=False, sheet_name="enriched")

    wb = load_workbook(path)
    ws = wb.active
    columns = list(safe.columns)
    col_index = {name: i + 1 for i, name in enumerate(columns)}

    _style_header(ws, len(columns))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _set_column_widths(ws, safe, columns)

    review_col = col_index.get("needs_review")
    conf_col = col_index.get("match_confidence")

    for i in range(len(safe)):
        excel_row = i + 2
        record = safe.iloc[i]

        if review_col and _is_true(record.get("needs_review")):
            for col in range(1, len(columns) + 1):
                ws.cell(row=excel_row, column=col).fill = _REVIEW_FILL

        if conf_col:
            style = _CONFIDENCE_STYLE.get(str(record.get("match_confidence")).strip().lower())
            if style:
                cell = ws.cell(row=excel_row, column=conf_col)
                cell.fill, cell.font = style

        for name in _URL_COLUMNS:
            idx = col_index.get(name)
            if not idx:
                continue
            url = str(record.get(name) or "").strip()
            if url.startswith("http"):
                cell = ws.cell(row=excel_row, column=idx)
                cell.hyperlink = url
                cell.font = _LINK_FONT

    wb.save(path)
    return path


def _style_header(ws, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _set_column_widths(ws, df: pd.DataFrame, columns) -> None:
    for i, name in enumerate(columns, start=1):
        widest = len(str(name))
        if len(df):
            widest = max(widest, int(df[name].astype(str).map(len).max()))
        width = max(_MIN_WIDTH, min(_MAX_WIDTH, widest + 2))
        ws.column_dimensions[get_column_letter(i)].width = width


def _is_true(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() == "true"
