"""Tests for src.excel_format (and io.write_table XLSX styling)."""
import pandas as pd
from openpyxl import load_workbook

from src import io as table_io
from src.excel_format import write_formatted_xlsx


def sample_df():
    return pd.DataFrame(
        [
            {
                "practice_name": "Acme Dental",
                "google_website": "https://acme.example",
                "google_maps_uri": "https://maps.google.com/?cid=1",
                "match_confidence": "low",
                "needs_review": True,
            },
            {
                "practice_name": "Beta Clinic",
                "google_website": "https://beta.example",
                "google_maps_uri": "",
                "match_confidence": "high",
                "needs_review": False,
            },
        ]
    )


def _cols(df):
    return {name: i + 1 for i, name in enumerate(df.columns)}


class TestFormatting:
    def test_freeze_filter_and_header(self, tmp_path):
        write_formatted_xlsx(sample_df(), tmp_path / "out.xlsx")
        ws = load_workbook(tmp_path / "out.xlsx").active
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref  # filters applied across the table
        assert ws.cell(row=1, column=1).font.bold

    def test_column_widths_set(self, tmp_path):
        write_formatted_xlsx(sample_df(), tmp_path / "out.xlsx")
        ws = load_workbook(tmp_path / "out.xlsx").active
        assert ws.column_dimensions["A"].width >= 10

    def test_needs_review_row_highlighted(self, tmp_path):
        df = sample_df()
        write_formatted_xlsx(df, tmp_path / "out.xlsx")
        ws = load_workbook(tmp_path / "out.xlsx").active
        # Row 2 is needs_review=True -> pink fill; row 3 is not.
        assert ws.cell(row=2, column=1).fill.fgColor.rgb.endswith("FCE4E4")
        assert not ws.cell(row=3, column=1).fill.fgColor.rgb.endswith("FCE4E4")

    def test_confidence_styled_distinctly(self, tmp_path):
        df = sample_df()
        write_formatted_xlsx(df, tmp_path / "out.xlsx")
        ws = load_workbook(tmp_path / "out.xlsx").active
        conf = _cols(df)["match_confidence"]
        low_cell = ws.cell(row=2, column=conf)
        high_cell = ws.cell(row=3, column=conf)
        assert low_cell.fill.fgColor.rgb.endswith("FFCC99")   # low = orange
        assert high_cell.fill.fgColor.rgb.endswith("C6EFCE")  # high = green
        assert low_cell.font.color.rgb != high_cell.font.color.rgb

    def test_urls_are_clickable(self, tmp_path):
        df = sample_df()
        write_formatted_xlsx(df, tmp_path / "out.xlsx")
        ws = load_workbook(tmp_path / "out.xlsx").active
        cols = _cols(df)
        web = ws.cell(row=2, column=cols["google_website"])
        assert web.hyperlink is not None
        assert web.hyperlink.target == "https://acme.example"
        # google_maps_uri present on row 2, blank on row 3.
        assert ws.cell(row=2, column=cols["google_maps_uri"]).hyperlink is not None
        assert ws.cell(row=3, column=cols["google_maps_uri"]).hyperlink is None

    def test_write_table_routes_xlsx_through_formatter(self, tmp_path):
        out = tmp_path / "via_io.xlsx"
        table_io.write_table(sample_df(), out)
        ws = load_workbook(out).active
        assert ws.freeze_panes == "A2"  # styling was applied

    def test_missing_columns_do_not_crash(self, tmp_path):
        # A DataFrame without the enrichment columns still writes fine.
        df = pd.DataFrame([{"a": 1, "b": 2}])
        write_formatted_xlsx(df, tmp_path / "plain.xlsx")
        ws = load_workbook(tmp_path / "plain.xlsx").active
        assert ws.cell(row=1, column=1).value == "a"
